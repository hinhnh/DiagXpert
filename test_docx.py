#!/usr/bin/env python3
"""
Script test tạo file DOCX thực sự
"""

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_test_docx():
    """Tạo file DOCX test đơn giản"""
    
    # Tạo document mới
    doc = Document()
    
    # Thêm tiêu đề
    title = doc.add_heading('HỆ THỐNG TRUYỀN ĐỘNG Ô TÔ', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Thêm đoạn văn
    doc.add_paragraph('Đây là file DOCX thực sự được tạo bởi python-docx.')
    
    # Thêm danh sách
    doc.add_heading('Các thành phần chính:', level=1)
    components = ['Hộp số', 'Ly hợp', 'Trục truyền', 'Bộ vi sai']
    for component in components:
        doc.add_paragraph(component, style='List Bullet')
    
    # Lưu file
    doc.save('resource/test_docx.docx')
    print("✅ Đã tạo file DOCX test: resource/test_docx.docx")

if __name__ == "__main__":
    create_test_docx()
