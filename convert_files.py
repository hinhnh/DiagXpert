#!/usr/bin/env python3
"""
Script convert các file text thành PDF, DOCX, DOC
"""

import os
from pathlib import Path

def create_pdf_from_text(text_file, pdf_file):
    """Tạo PDF từ text file"""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        
        # Đọc nội dung text
        with open(text_file, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # Tạo PDF
        doc = SimpleDocTemplate(pdf_file, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        
        # Chia nội dung thành paragraphs
        paragraphs = content.split('\n\n')
        for para in paragraphs:
            if para.strip():
                # Xử lý tiêu đề
                if para.startswith('HỆ THỐNG') or para.startswith('TỔNG QUAN') or para.startswith('CÁC THÀNH PHẦN'):
                    story.append(Paragraph(para, styles['Heading1']))
                elif para.startswith('1.') or para.startswith('2.') or para.startswith('3.') or para.startswith('4.'):
                    story.append(Paragraph(para, styles['Heading2']))
                else:
                    story.append(Paragraph(para, styles['Normal']))
                story.append(Spacer(1, 12))
        
        doc.build(story)
        print(f"✅ Đã tạo PDF: {pdf_file}")
        return True
        
    except ImportError:
        print(f"⚠️ Không thể tạo PDF - cần cài reportlab")
        return False
    except Exception as e:
        print(f"❌ Lỗi tạo PDF: {e}")
        return False

def create_docx_from_text(text_file, docx_file):
    """Tạo DOCX từ text file"""
    try:
        from docx import Document
        from docx.shared import Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        # Đọc nội dung text
        with open(text_file, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # Tạo DOCX
        doc = Document()
        
        # Chia nội dung thành paragraphs
        paragraphs = content.split('\n\n')
        for para in paragraphs:
            if para.strip():
                # Xử lý tiêu đề
                if para.startswith('HỆ THỐNG') or para.startswith('TỔNG QUAN'):
                    doc.add_heading(para, level=1)
                elif para.startswith('1.') or para.startswith('2.') or para.startswith('3.') or para.startswith('4.'):
                    doc.add_heading(para, level=2)
                elif para.startswith('CÁC THÀNH PHẦN') or para.startswith('CÁC LOẠI') or para.startswith('CHẨN ĐOÁN') or para.startswith('BẢO DƯỠNG'):
                    doc.add_heading(para, level=2)
                else:
                    doc.add_paragraph(para)
        
        # Lưu file
        doc.save(str(docx_file))
        print(f"✅ Đã tạo DOCX: {docx_file}")
        return True
        
    except ImportError:
        print(f"⚠️ Không thể tạo DOCX - cần cài python-docx")
        return False
    except Exception as e:
        print(f"❌ Lỗi tạo DOCX: {e}")
        return False

def create_doc_from_text(text_file, doc_file):
    """Tạo DOC từ text file (sử dụng python-docx)"""
    try:
        from docx import Document
        
        # Đọc nội dung text
        with open(text_file, 'r', encoding='utf-8') as f:
            content = f.read()
        
        # Tạo DOC (thực ra là DOCX nhưng đổi tên)
        doc = Document()
        
        # Chia nội dung thành paragraphs
        paragraphs = content.split('\n\n')
        for para in paragraphs:
            if para.strip():
                # Xử lý tiêu đề tương tự DOCX
                if para.startswith('HỆ THỐNG') or para.startswith('TỔNG QUAN'):
                    doc.add_heading(para, level=1)
                elif para.startswith('1.') or para.startswith('2.') or para.startswith('3.') or para.startswith('4.'):
                    doc.add_heading(para, level=2)
                elif para.startswith('CÁC THÀNH PHẦN') or para.startswith('CÁC LOẠI') or para.startswith('CHẨN ĐOÁN') or para.startswith('BẢO DƯỠNG'):
                    doc.add_heading(para, level=2)
                else:
                    doc.add_paragraph(para)
        
        # Lưu với tên .doc (thực ra vẫn là DOCX format)
        doc.save(str(doc_file))
        print(f"✅ Đã tạo DOC: {doc_file}")
        return True
        
    except ImportError:
        print(f"⚠️ Không thể tạo DOC - cần cài python-docx")
        return False
    except Exception as e:
        print(f"❌ Lỗi tạo DOC: {e}")
        return False

def main():
    """Convert tất cả file"""
    resource_dir = Path("resource")
    
    if not resource_dir.exists():
        print("❌ Không tìm thấy thư mục resource")
        return
    
    print("🔄 Đang convert các file...")
    
    # Convert từ TXT sang các định dạng khác
    text_files = [
        "air_conditioning_system.txt",
        "fuel_injection_system.pdf",  # Thực ra là .txt
        "transmission_system.docx",   # Thực ra là .txt
        "suspension_system.doc"       # Thực ra là .txt
    ]
    
    for text_file in text_files:
        text_path = resource_dir / text_file
        if text_path.exists():
            base_name = text_path.stem
            
            # Convert sang PDF
            pdf_file = resource_dir / f"{base_name}.pdf"
            if not pdf_file.exists():
                create_pdf_from_text(text_path, str(pdf_file))
            
            # Convert sang DOCX
            docx_file = resource_dir / f"{base_name}.docx"
            if not docx_file.exists():
                create_docx_from_text(text_path, docx_file)
            
            # Convert sang DOC
            doc_file = resource_dir / f"{base_name}.doc"
            if not doc_file.exists():
                create_doc_from_text(text_path, doc_file)
    
    print("\n📁 Kiểm tra thư mục resource:")
    for file in resource_dir.glob("*"):
        print(f"  📄 {file.name}")
    
    print("\n✅ Hoàn thành convert files!")

if __name__ == "__main__":
    main()
